# ABOUTME: Tests μηνιαίου υπολογισμού ΦΠΑ (εκροές/εισροές/εκπιπτόμενο/υπόλοιπο).
# ABOUTME: Καλύπτει εκπιπτόμενο ΦΠΑ, πράξεις λήπτη (flag) και υπόλοιπο.
from datetime import date
import pytest
from esoda_exoda.models import SheetRow
from esoda_exoda import vat
from esoda_exoda.vat import compute

def _inc(m, net, vat):
    return SheetRow("2.1","","",date(2025,m,10),"990000106",net,vat,net+vat,0.0,"ΝΑΙ","c","k")
def _exp(m, net, vat, ekpiptei, kind="Λοιπά έξοδα (E3_585_016)"):
    return SheetRow("1.1","","",date(2025,m,5),"090000125",net,vat,net+vat,ekpiptei,"","cat",kind)

def test_monthly_balance_and_deductible():
    v = compute([_inc(1, 1000, 240)], [_exp(1, 100, 24, "ΝΑΙ"), _exp(1, 50, 12, "ΌΧΙ")])
    assert len(v) == 1
    m = v[0]
    assert m.month == 1
    assert round(m.out_vat, 2) == 240.0
    assert round(m.in_vat, 2) == 36.0
    assert round(m.in_vat_deductible, 2) == 24.0        # μόνο το ΝΑΙ
    assert round(m.balance, 2) == 216.0                  # 240 - 24

def test_reverse_charge_flagged_not_double_counted():
    v = compute([], [_exp(3, 200, 0, "ΝΑΙ", "Λοιπές πράξεις λήπτη (κωδ.366-Φ2)")])
    assert v[0].reverse_charge_net == 200.0
    assert v[0].out_vat == 0.0                            # δεν αυτοχρεώνεται (απλή σύνοψη)

def test_reverse_charge_shows_both_legs_and_zero_net_effect():
    r143 = SheetRow("14.3 Τιμολόγιο / Ενδοκοινοτική Λήψη Υπηρεσιών", "0", "1",
                    date(2026, 7, 31), "NL999999003B01", 100.0, 24.0, 124.0,
                    "ΝΑΙ", "ΝΑΙ", "2.3 Λήψη Υπηρεσιών",
                    "Λοιπές αμοιβές για υπηρεσίες αλλοδαπής (E3_585_010)")
    r144 = SheetRow("14.4 Τιμολόγιο / Λήψη Υπηρεσιών Τρίτων Χωρών", "0", "1",
                    date(2026, 7, 15), "000000000", 50.0, 12.0, 62.0,
                    "ΝΑΙ", "ΝΑΙ", "2.3 Λήψη Υπηρεσιών",
                    "Λοιπές αμοιβές για υπηρεσίες αλλοδαπής (E3_585_010)")
    months = compute([], [r143, r144])
    assert len(months) == 1
    m = months[0]
    assert m.rc_net_365 == 100.0 and m.rc_net_366 == 50.0
    assert m.rc_vat == 36.0                  # χρέωση (σκέλος εκροών)
    assert m.in_vat_deductible == 36.0       # έκπτωση (σκέλος εισροών)
    assert m.balance == 0.0                  # μηδενικό καθαρό αποτέλεσμα

def test_non_14x_rows_do_not_touch_reverse_charge_fields():
    exp = SheetRow("1.1 Τιμολόγιο", "", "", date(2026, 7, 1), "090000125",
                   100.0, 24.0, 124.0, "ΝΑΙ", "ΝΑΙ", "2.3", "Λοιπά (E3_585_016)")
    m = compute([], [exp])[0]
    assert m.rc_net_365 == 0.0 and m.rc_net_366 == 0.0 and m.rc_vat == 0.0

def test_άγνωστο_δεν_μετράει_στο_εκπιπτόμενο():
    rows = [SheetRow(type="2.1", series="", aa="1", date=date(2026, 7, 1), afm="1",
                     net=100.0, vat=24.0, total=124.0, col_j="ΑΓΝΩΣΤΟ",
                     charact="", category="", kind="")]
    vm = compute([], rows)[0]
    assert vm.in_vat == 24.0
    assert vm.in_vat_deductible == 0.0

def test_in_vat_είναι_εκπιπτόμενο_συν_άγνωστα_συν_ρητό_όχι():
    # Invariant: κάθε ευρώ ΦΠΑ εισροών πρέπει να καταλήγει ΑΚΡΙΒΩΣ σε μία από τις τρεις
    # μοίρες — εκπίπτει, άγνωστο, ρητά δεν εκπίπτει — αλλιώς εξαφανίζεται σιωπηλά.
    rows = [_exp(7, 100, 10, "ΝΑΙ"), _exp(7, 100, 5, "ΌΧΙ"),
           _exp(7, 100, 3, "ΑΓΝΩΣΤΟ"), _exp(7, 100, 2, None)]
    vm = compute([], rows)[0]
    explicit_no_vat = sum(r.vat for r in rows if str(r.col_j or "").strip().upper()
                          in ("ΌΧΙ", "ΟΧΙ", "OXI"))
    assert vm.in_vat == pytest.approx(20.0)
    assert explicit_no_vat == pytest.approx(5.0)
    assert vm.in_vat == pytest.approx(
        vm.in_vat_deductible + vm.unknown_vat + explicit_no_vat)

def test_κενό_col_j_μετράει_στο_άγνωστα():
    for empty in (None, ""):
        rows = [_exp(7, 100, 7, empty)]
        vm = compute([], rows)[0]
        assert vm.unknown_vat == pytest.approx(7.0)
        assert vm.in_vat_deductible == 0.0

def test_compute_γεμίζει_τα_επίσημα_και_τη_διαφορά():
    income = [SheetRow(type="1.1", series="", aa="1", date=date(2026, 7, 10), afm="1",
                       net=1250.0, vat=300.0, total=1550.0, col_j=0.0,
                       charact="ΝΑΙ", category="", kind="")]
    expenses = [SheetRow(type="2.1", series="", aa="1", date=date(2026, 7, 8), afm="2",
                         net=50.0, vat=12.0, total=62.0, col_j="ΌΧΙ",
                         charact="ΝΑΙ", category="", kind=""),
                SheetRow(type="2.1", series="", aa="2", date=date(2026, 7, 22), afm="3",
                         net=115.0, vat=27.6, total=142.6, col_j="ΝΑΙ",
                         charact="ΝΑΙ", category="", kind="")]
    official = {7: {"Vat333": 327.6, "Vat386": 27.6}}
    vm = vat.compute(income, expenses, official)[0]
    assert vm.out_vat == 300.0
    assert vm.in_vat_deductible == 27.6
    assert vm.balance == pytest.approx(272.4)
    assert vm.official_out_vat == 327.6
    assert vm.official_deductible == 27.6
    assert vm.official_balance == pytest.approx(300.0)
    # Η υπόλοιπη διαφορά είναι η αυτο-χρέωση εκροών των πράξεων λήπτη, που τα βιβλία
    # δεν κάνουν. Το plan την κάνει ορατή, δεν την κλείνει.
    assert vm.diff == pytest.approx(-27.6)

def test_compute_το_381_είναι_το_κανονικό_επίσημο_εκπιπτόμενο():
    # Μάρτιος 2026: κανονικές αγορές εσωτερικού (VAT_361) δίνουν 361/381, χωρίς κανένα 386.
    official = {3: {"Vat303": 2500.0, "Vat333": 600.0, "Vat361": 25.0, "Vat381": 6.0}}
    vm = compute([_inc(3, 2500.0, 600.0)], [_exp(3, 25.0, 6.0, "ΝΑΙ")], official)[0]
    assert vm.official_deductible == pytest.approx(6.0)
    assert vm.diff == pytest.approx(0.0)

def test_compute_το_επίσημο_εκπιπτόμενο_αθροίζει_όλα_τα_πεδία_381_έως_386():
    official = {7: {"Vat381": 1.0, "Vat382": 2.0, "Vat383": 4.0, "Vat384": 8.0,
                    "Vat385": 16.0, "Vat386": 32.0}}
    vm = compute([], [], official)[0]
    assert vm.official_deductible == pytest.approx(63.0)

def test_unknown_codes_δεν_αναφέρει_τα_πεδία_εισροών():
    inputs = {f"Vat{c}": 1.0 for c in (*range(361, 367), *range(381, 387))}
    assert vat.unknown_codes({1: {**inputs, "Vat331": 5.0}}) == {1: {"Vat331": 5.0}}
    assert vat.unknown_codes({1: inputs}) == {}

def test_compute_μετρά_τα_άγνωστα_χωριστά():
    expenses = [SheetRow(type="2.1", series="", aa="1", date=date(2026, 7, 8), afm="2",
                         net=100.0, vat=24.0, total=124.0, col_j="ΑΓΝΩΣΤΟ",
                         charact="", category="", kind="")]
    vm = vat.compute([], expenses)[0]
    assert vm.unknown_vat == 24.0
    assert vm.in_vat_deductible == 0.0

def test_compute_χωρίς_επίσημα_δίνει_μηδέν_και_δεν_σκάει():
    expenses = [SheetRow(type="2.1", series="", aa="1", date=date(2026, 7, 8), afm="2",
                         net=100.0, vat=24.0, total=124.0, col_j="ΝΑΙ",
                         charact="", category="", kind="")]
    vm = vat.compute([], expenses)[0]
    assert (vm.official_out_vat, vm.official_deductible, vm.official_balance) == (0.0, 0.0, 0.0)
    assert vm.diff == pytest.approx(-24.0)

def test_write_vat_sheet_γράφει_τις_νέες_στήλες():
    # Δύο μη μηδενικές, διακριτές τιμές στις στήλες 12 (Άγνωστα) και 16 (ΔΙΑΦΟΡΑ):
    # αν οι στήλες ήταν ανταλλαγμένες, ή η 16 τρεφόταν λάθος από το unknown_vat, το
    # test θα το έπιανε — με δύο μηδενικά (όπως πριν) δεν το έπιανε τίποτα.
    import openpyxl
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    official = {7: {"Vat333": 327.6, "Vat386": 27.6}}
    income = [SheetRow(type="1.1", series="", aa="1", date=date(2026, 7, 10), afm="1",
                       net=1250.0, vat=300.0, total=1550.0, col_j=0.0,
                       charact="", category="", kind="")]
    expenses = [SheetRow(type="2.1", series="", aa="1", date=date(2026, 7, 8), afm="2",
                         net=50.0, vat=12.0, total=62.0, col_j="ΌΧΙ",
                         charact="ΝΑΙ", category="", kind=""),
                SheetRow(type="2.1", series="", aa="2", date=date(2026, 7, 22), afm="3",
                         net=115.0, vat=27.6, total=142.6, col_j="ΝΑΙ",
                         charact="ΝΑΙ", category="", kind=""),
                SheetRow(type="2.1", series="", aa="3", date=date(2026, 7, 15), afm="4",
                         net=20.0, vat=5.0, total=25.0, col_j="ΑΓΝΩΣΤΟ",
                         charact="", category="", kind="")]
    months = vat.compute(income, expenses, official)
    vat.write_vat_sheet(wb, months)
    ws = wb["ΦΠΑ"]
    headers = [c.value for c in ws[1]]
    assert headers[11:] == ["Άγνωστα (ΦΠΑ)", "Επίσημο ΦΠΑ εκροών (333)",
                            "Επίσημο εκπιπτόμενο (381–386)", "Επίσημο υπόλοιπο",
                            "ΔΙΑΦΟΡΑ (βιβλία − επίσημο)"]
    assert ws.cell(row=2, column=13).value == 327.6
    # Στήλη 12: μόνο η γραμμή ΑΓΝΩΣΤΟ (5.0) — οι ΝΑΙ/ΌΧΙ δεν συνεισφέρουν εδώ.
    assert ws.cell(row=2, column=12).value == pytest.approx(5.0)
    # Στήλη 16: balance (300.0 - 27.6 = 272.4) − official_balance
    # (327.6 - 27.6 = 300.0) = -27.6 — η αυτο-χρέωση εκροών των πράξεων λήπτη
    # που τα βιβλία δεν κάνουν (§Global Constraints). Μη μηδενική, διαφορετική από
    # τη στήλη 12 (5.0).
    assert ws.cell(row=2, column=16).value == pytest.approx(-27.6)
    assert ws.cell(row=2, column=12).value != ws.cell(row=2, column=16).value
