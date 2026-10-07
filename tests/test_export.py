# ABOUTME: Tests του Excel export από τη βάση: τα φύλλα, οι τιμές και τα φίλτρα
# ABOUTME: (έτος/μήνας) του αρχείου, χωρίς κανένα αρχείο στο δίσκο.
import io
from datetime import date
import openpyxl
from esoda_exoda.counterparties import Counterparty, ensure_counterparty, save_counterparty, save_default
from esoda_exoda.export import build_workbook
from esoda_exoda.fetches import store_fetch
from esoda_exoda.invoices import add_snapshot, upsert_book_record
from esoda_exoda.models import Record

def _seed(db, taxpayer):
    f = store_fetch(db, taxpayer.id, "RequestMyIncome", date(2026, 1, 1), date(2026, 12, 31), "")
    save_counterparty(db, Counterparty(id=None, taxpayer_id=taxpayer.id, vat="990000106", name="ΠΕΛΑΤΗΣ Α", separate_totals=True))
    elidek = ensure_counterparty(db, taxpayer.id, "990000106")
    save_default(db, elidek, "income", "category1_3", "E3_561_001", None)
    upsert_book_record(db, taxpayer.id, Record("990000106", date(2026, 7, 1), "2.1", 500, 120, 100, 620, "I1"), "income", f, elidek)
    sup = ensure_counterparty(db, taxpayer.id, "090000113", "Π ΟΕ")
    save_default(db, sup, "expense", "category2_4", "E3_585_016", True)
    upsert_book_record(db, taxpayer.id, Record("090000113", date(2026, 7, 3), "1.1", 100, 24, 0, 124, "M1"), "expense", f, sup)
    upsert_book_record(db, taxpayer.id, Record("1", date(2026, 8, 3), "1.1", 10, 2.4, 0, 12.4, "M2"), "expense", f, None)
    add_snapshot(db, taxpayer.id, "I1", "vat_info", f, date(2026, 7, 1), codes={"Vat333": 120.0, "Vat386": 24.0, "Vat999": 1.0})

def test_export_έχει_τα_τέσσερα_φύλλα_και_τη_σύνοψη(db, db_cfg, taxpayer):
    _seed(db, taxpayer)
    data, summary = build_workbook(db, db_cfg, taxpayer, 2026)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["ΕΣΟΔΑ", "ΕΞΟΔΑ", "Ε3", "ΦΠΑ"]
    assert summary["ΕΣΟΔΑ"] == {"rows": 1, "net": 500.0, "vat": 120.0}
    assert summary["ΕΞΟΔΑ"]["rows"] == 2
    july = next(v for v in summary["ΦΠΑ"] if v["month"] == 7)
    assert july["in_vat_deductible"] == 24.0 and july["official_deductible"] == 24.0
    assert summary["ΦΠΑ_λοιποί_κωδικοί"] == {7: {"Vat999": 1.0}}
    assert wb["ΕΣΟΔΑ"].cell(row=5, column=1).value == "ΣΥΝΟΛΑ ΠΕΛΑΤΗΣ Α (ΑΦΜ 990000106)"
    assert wb["ΕΞΟΔΑ"].cell(row=3, column=10).value == "ΑΓΝΩΣΤΟ"

def test_export_μήνα_φιλτράρει(db, db_cfg, taxpayer):
    _seed(db, taxpayer)
    data, summary = build_workbook(db, db_cfg, taxpayer, 2026, 8)
    assert summary["ΕΞΟΔΑ"]["rows"] == 1 and summary["ΕΣΟΔΑ"]["rows"] == 0
