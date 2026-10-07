# ABOUTME: Tests του φύλλου Ε3: κάθε κωδικός εξόδων έχει γραμμή, τα αχαρακτήριστα φαίνονται,
# ABOUTME: και η γραμμή ελέγχου πιάνει κάθε ποσό που δεν ανήκει σε καμία γραμμή.
import openpyxl
from esoda_exoda.e3 import write_e3_sheet

G = "ΕΞΟΔΑ!$G$2:$G$6"; H = "ΕΞΟΔΑ!$H$2:$H$6"; J = "ΕΞΟΔΑ!$J$2:$J$6"; M = "ΕΞΟΔΑ!$M$2:$M$6"

def _sheet():
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    write_e3_sheet(wb, n_income=3, n_expenses=5)
    return wb["Ε3"]

def _row(ws, desc_prefix):
    [r] = [r for r in range(1, ws.max_row + 1)
           if str(ws.cell(row=r, column=2).value or "").startswith(desc_prefix)]
    return r

def test_585_013_τηλεπικοινωνίες_έχει_γραμμή():
    ws = _sheet(); r = _row(ws, "Τηλεπικοινωνίες")
    assert ws.cell(row=r, column=1).value == "585_013"
    assert ws.cell(row=r, column=3).value == (f'=SUMIFS({G},{M},"*E3_585_013*")'
                                              f'+SUMIFS({H},{M},"*E3_585_013*",{J},"ΌΧΙ")')

def test_τα_έξοδα_χωρίς_κωδικό_ε3_φαίνονται():
    ws = _sheet(); r = _row(ws, "Χωρίς κωδικό Ε3")
    assert ws.cell(row=r, column=3).value == (f'=SUMIFS({G},{M},"")'
                                              f'+SUMIFS({H},{M},"",{J},"ΌΧΙ")')

def test_ο_έλεγχος_αφαιρεί_από_το_σύνολο_κάθε_γραμμή_εξόδων():
    ws = _sheet()
    first = next(r for r in range(1, ws.max_row + 1)
                 if str(ws.cell(row=r, column=1).value or "").startswith("ΕΞΟΔΑ"))
    check = _row(ws, "Λοιποί κωδικοί Ε3")
    amount_rows = [r for r in range(first, check)
                   if str(ws.cell(row=r, column=3).value or "").startswith("=")]
    assert len(amount_rows) >= 9
    assert ws.cell(row=check, column=3).value == (
        f'=SUM({G})+SUMIFS({H},{J},"ΌΧΙ")' + "".join(f"-C{r}" for r in amount_rows))
