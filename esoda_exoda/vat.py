# ABOUTME: Υπολογισμός μηνιαίας σύνοψης ΦΠΑ από γραμμές ΕΣΟΔΩΝ/ΕΞΟΔΩΝ.
# ABOUTME: Οι πράξεις λήπτη 14.3/14.4 χρεώνονται και εκπίπτονται ισόποσα στο υπόλοιπο.
from dataclasses import dataclass
from collections import defaultdict
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

@dataclass
class VatMonth:
    month: int
    out_net: float = 0.0
    out_vat: float = 0.0
    in_net: float = 0.0
    in_vat: float = 0.0
    in_vat_deductible: float = 0.0
    balance: float = 0.0
    reverse_charge_net: float = 0.0
    rc_net_365: float = 0.0
    rc_net_366: float = 0.0
    rc_vat: float = 0.0
    unknown_vat: float = 0.0
    official_out_vat: float = 0.0
    official_deductible: float = 0.0
    official_balance: float = 0.0
    diff: float = 0.0

def is_reverse_charge(row) -> bool:
    k = row.kind or ""
    return "366" in k or "ληπτ" in k.casefold()

def _yes(v) -> bool:
    return str(v or "").strip().upper() in ("ΝΑΙ", "NAI")

_EXPLICIT_NO = ("ΌΧΙ", "ΟΧΙ", "OXI")

def _unknown(v) -> bool:
    """Ό,τι δεν είναι ρητό ΝΑΙ ούτε ρητό ΌΧΙ είναι απουσία απόφασης, όχι απόφαση: κενό
    κελί από χειροκίνητη επεξεργασία, ή τιμή με λάθος τόνο. Πρέπει να φαίνεται στη στήλη
    «Άγνωστα», αλλιώς το ΦΠΑ του εξαφανίζεται και από τις δύο στήλες χωρίς ίχνος."""
    return not _yes(v) and str(v or "").strip().upper() not in _EXPLICIT_NO

# Στις εκροές το «επίσημο υπόλοιπο» παίρνει μόνο το κλιμάκιο 24% (303/333), το μόνο
# που εμφανίζεται στα πραγματικά δεδομένα του 2026. Στις εισροές, τα πεδία 381–386
# είναι το ΦΠΑ των 361–366 ανά είδος εισροής, όχι ανά συντελεστή (τεκμηρίωση API
# myDATA): 381 οι αγορές εσωτερικού (VAT_361), 386 οι λοιπές πράξεις λήπτη.
OFFICIAL_OUT_VAT = "Vat333"
OFFICIAL_DEDUCTIBLE = tuple(f"Vat{c}" for c in range(381, 387))
# Οι κωδικοί που ΑΝΑΓΝΩΡΙΖΟΥΜΕ: 333 και 381–386 μπαίνουν στον τύπο, 303 και 361–366
# είναι καθαρές αξίες και αγνοούνται σκόπιμα. Ό,τι δεν είναι εδώ, αναφέρεται — ένα
# κλιμάκιο 13%/6% εκροών πρέπει να φαίνεται, όχι να απορροφηθεί σιωπηλά σε νούμερο δήλωσης.
_KNOWN_CODES = (OFFICIAL_OUT_VAT, *OFFICIAL_DEDUCTIBLE, "Vat303",
                *(f"Vat{c}" for c in range(361, 367)))

def unknown_codes(official: dict[int, dict[str, float]]) -> dict[int, dict[str, float]]:
    """Κωδικοί Φ2 που ο τύπος του επίσημου υπολοίπου δεν αναγνωρίζει."""
    return {m: {c: v for c, v in codes.items() if c not in _KNOWN_CODES}
            for m, codes in official.items()
            if any(c not in _KNOWN_CODES for c in codes)}

def compute(income_rows, expense_rows, official=None):
    official = official or {}
    months = defaultdict(lambda: VatMonth(0))
    for r in income_rows:
        m = months[r.date.month]; m.month = r.date.month
        m.out_net += r.net; m.out_vat += r.vat
    for r in expense_rows:
        m = months[r.date.month]; m.month = r.date.month
        m.in_net += r.net; m.in_vat += r.vat
        if _yes(r.col_j):
            m.in_vat_deductible += r.vat
        elif _unknown(r.col_j):
            m.unknown_vat += r.vat
        if is_reverse_charge(r):
            m.reverse_charge_net += r.net
        if r.type.startswith("14.3"):
            m.rc_net_365 += r.net; m.rc_vat += r.vat
        elif r.type.startswith("14.4"):
            m.rc_net_366 += r.net; m.rc_vat += r.vat
    # Μήνας με επίσημα στοιχεία αλλά χωρίς γραμμές βιβλίων πρέπει να εμφανιστεί στο
    # φύλλο: αλλιώς μια ολόκληρη ΔΙΑΦΟΡΑ εξαφανίζεται μαζί με τη γραμμή.
    for month in official:
        months[month].month = month
    out = []
    for month in sorted(months):
        vm = months[month]
        vm.balance = vm.out_vat + vm.rc_vat - vm.in_vat_deductible
        codes = official.get(month, {})
        vm.official_out_vat = codes.get(OFFICIAL_OUT_VAT, 0.0)
        vm.official_deductible = sum(codes.get(c, 0.0) for c in OFFICIAL_DEDUCTIBLE)
        vm.official_balance = vm.official_out_vat - vm.official_deductible
        vm.diff = vm.balance - vm.official_balance
        out.append(vm)
    return out

HEADERS = ["Μήνας", "Καθ.Αξία Εκροών", "ΦΠΑ Εκροών", "Καθ.Αξία Εισροών",
           "ΦΠΑ Εισροών", "Εκπιπτόμενο ΦΠΑ", "Υπόλοιπο", "Πράξεις λήπτη (καθ.αξία)",
           "Λήπτη 365 (καθ.)", "Λήπτη 366 (καθ.)", "ΦΠΑ λήπτη (χρέωση=έκπτωση)",
           "Άγνωστα (ΦΠΑ)", "Επίσημο ΦΠΑ εκροών (333)", "Επίσημο εκπιπτόμενο (381–386)",
           "Επίσημο υπόλοιπο", "ΔΙΑΦΟΡΑ (βιβλία − επίσημο)"]
WIDTHS = [8, 16, 14, 16, 14, 18, 14, 24, 16, 16, 24, 14, 24, 24, 18, 26]

def write_vat_sheet(wb, vat_months: list[VatMonth]) -> None:
    ws = wb.create_sheet("ΦΠΑ")
    bold = Font(bold=True)
    for col, h in enumerate(HEADERS, start=1):
        ws.cell(row=1, column=col, value=h).font = bold
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    r = 2
    for vm in vat_months:
        ws.cell(row=r, column=1, value=vm.month)
        for col, val in ((2, vm.out_net), (3, vm.out_vat), (4, vm.in_net),
                         (5, vm.in_vat), (6, vm.in_vat_deductible),
                         (7, vm.balance), (8, vm.reverse_charge_net),
                         (9, vm.rc_net_365), (10, vm.rc_net_366), (11, vm.rc_vat),
                         (12, vm.unknown_vat), (13, vm.official_out_vat),
                         (14, vm.official_deductible), (15, vm.official_balance),
                         (16, vm.diff)):
            c = ws.cell(row=r, column=col, value=val); c.number_format = "#,##0.00"
        r += 1
    first, last = 2, r - 1
    ws.cell(row=r, column=1, value="ΣΥΝΟΛΑ").font = bold
    for col in range(2, 17):
        L = get_column_letter(col)
        c = ws.cell(row=r, column=col, value=f"=SUM({L}{first}:{L}{last})")
        c.number_format = "#,##0.00"; c.font = bold
