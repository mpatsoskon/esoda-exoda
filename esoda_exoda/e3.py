# ABOUTME: Γράφει το φύλλο Ε3 με τύπους πάνω στα φύλλα ΕΣΟΔΑ/ΕΞΟΔΑ του ΣΥΝΟΛΑ.
# ABOUTME: Κάθε κωδικός Ε3 αθροίζει καθαρή αξία + μη εκπιπτόμενο ΦΠΑ των σχετικών γραμμών.
from openpyxl.styles import Font

def write_e3_sheet(wb, n_income: int, n_expenses: int) -> None:
    ws = wb.create_sheet("Ε3")
    bold = Font(bold=True)
    inc_total_row = max(2, n_income + 1) + 2  # ίδια αριθμητική με το write_sheet (last=max(first, r-1))
    lo = 2
    hi = max(lo, n_expenses + 1)  # με μηδέν έξοδα το range μένει ορθό, όχι $G$2:$G$1
    G = f"ΕΞΟΔΑ!$G${lo}:$G${hi}"; H = f"ΕΞΟΔΑ!$H${lo}:$H${hi}"
    J = f"ΕΞΟΔΑ!$J${lo}:$J${hi}"; M = f"ΕΞΟΔΑ!$M${lo}:$M${hi}"
    def kind(code):
        return (f'=SUMIFS({G},{M},"*{code}*")'
                f'+SUMIFS({H},{M},"*{code}*",{J},"ΌΧΙ")')
    for col, w in zip("ABCD", (12, 52, 14, 62)):
        ws.column_dimensions[col].width = w
    for col, h in enumerate(("Κωδ. Ε3", "Περιγραφή", "Ποσό", "Σημείωση"), start=1):
        ws.cell(row=1, column=col, value=h).font = bold
    head = [
        ("ΕΣΟΔΑ (Πίνακας Ζ1)", None, None),
        ("561", "Πωλήσεις αγαθών και υπηρεσιών (E3_561_001)",
         f"=ΕΣΟΔΑ!G{inc_total_row}", None),
        (None, None, None),
    ]
    expenses = [
        ("ΕΞΟΔΑ (Πίνακας Ζ2 — λειτουργικά)", None, None),
        ("585_006", "Έξοδα ταξιδιών", kind("E3_585_006"), None),
        ("585_009", "Λοιπές αμοιβές υπηρεσιών ημεδαπής", kind("E3_585_009"), None),
        ("585_010", "Λοιπές αμοιβές για υπηρεσίες αλλοδαπής", kind("E3_585_010"), None),
        ("585_013", "Τηλεπικοινωνίες", kind("E3_585_013"), None),
        ("585_016", "Λοιπά έξοδα", kind("E3_585_016"), None),
        (None, None, None),
        ("ΠΑΓΙΑ", None, None),
        ("882", "Αγορές ενσώματων παγίων", kind("E3_882"), None),
        ("883", "Αγορές μη ενσώματων παγίων", kind("E3_883"), None),
        ("587", "Αποσβέσεις", kind("E3_587"), None),
        (None, "Χωρίς κωδικό Ε3 (να χαρακτηριστούν)",
         f'=SUMIFS({G},{M},"")+SUMIFS({H},{M},"",{J},"ΌΧΙ")', None),
    ]
    # Ένας κωδικός χωρίς δική του γραμμή εδώ δεν θα εμφανιζόταν πουθενά: ό,τι μένει από
    # το σύνολο των εξόδων αφού αφαιρεθούν όλες οι γραμμές είναι ακριβώς αυτά τα ποσά.
    first = 2 + len(head)
    subtract = "".join(f"-C{first + i}" for i, row in enumerate(expenses) if row[2])
    rows = head + expenses + [
        (None, "Λοιποί κωδικοί Ε3 (έλεγχος — πρέπει να είναι 0)",
         f'=SUM({G})+SUMIFS({H},{J},"ΌΧΙ")' + subtract, None),
        (None, "Παρακράτηση φόρου 20% (πάει στο Ε1)", f"=ΕΣΟΔΑ!J{inc_total_row}", None),
    ]
    for r, (code, desc, amount, *note) in enumerate(rows, start=2):
        if code and desc is None:
            ws.cell(row=r, column=1, value=code).font = bold; continue
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=desc)
        if amount is not None:
            c = ws.cell(row=r, column=3, value=amount)
            c.number_format = "#,##0.00"; c.font = bold
