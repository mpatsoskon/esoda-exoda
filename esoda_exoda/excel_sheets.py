# ABOUTME: Low-level εγγραφή/ανάγνωση φύλλων ΕΣΟΔΑ/ΕΞΟΔΑ με το σταθερό layout
# ABOUTME: (τιμές ανά γραμμή, τύποι SUM/SUMIF μόνο στις γραμμές συνόλων).
from datetime import datetime
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from .models import SheetRow

HEADERS = {
    "ΕΣΟΔΑ": ["", "", "Σειρά Παραστατικού", "ΑΑ ΠΑΡΑΣΤΑΤΙΚΟΥ", "ΗΜ/ΝΙΑ ΕΚΔΟΣΗΣ",
              "ΑΦΜ", "ΚΑΘ. ΑΞΙΑ", "ΦΠΑ", "ΣΥΝΟΛΟ", "ΦΟΡΟΣ",
              "ΧΑΡΑΚΤΗΡΙΣΜΟΣ ΕΣΟΔΟΥ", "Κατηγορία ΕΣΟΔΟΥ", "Είδος ΕΣΟΔΟΥ"],
    "ΕΞΟΔΑ": ["", "", "Σειρά Παραστατικού", "ΑΑ ΠΑΡΑΣΤΑΤΙΚΟΥ", "ΗΜ/ΝΙΑ ΕΚΔΟΣΗΣ",
              "ΑΦΜ", "ΚΑΘ. ΑΞΙΑ", "ΦΠΑ", "ΣΥΝΟΛΟ", "ΕΚΠΙΠΤΕΙ ΤΟ ΦΠΑ",
              "ΧΑΡΑΚΤΗΡΙΣΜΟΣ ΕΞΟΔΟΥ", "Κατηγορία ΕΞΟΔΟΥ", "Είδος ΕΞΟΔΟΥ"],
}
WIDTHS = [8, 30, 18, 16, 14, 12, 11, 11, 11, 15, 22, 30, 34]

def write_sheet(wb, name: str, rows: list[SheetRow], subtotal_parties=()) -> int:
    ws = wb.create_sheet(name)
    bold = Font(bold=True)
    for col, h in enumerate(HEADERS[name], start=1):
        ws.cell(row=1, column=col, value=h or None).font = bold
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    r = 2
    for n, row in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=2, value=row.type)
        ws.cell(row=r, column=3, value=row.series or None)
        ws.cell(row=r, column=4, value=row.aa or None)
        d = ws.cell(row=r, column=5, value=row.date); d.number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=6, value=row.afm or None)
        for col, val in ((7, row.net), (8, row.vat), (9, row.total)):
            c = ws.cell(row=r, column=col, value=val); c.number_format = "#,##0.00"
        if name == "ΕΣΟΔΑ":
            c = ws.cell(row=r, column=10, value=float(row.col_j or 0)); c.number_format = "#,##0.00"
        else:
            ws.cell(row=r, column=10, value=row.col_j)
        ws.cell(row=r, column=11, value=row.charact)
        ws.cell(row=r, column=12, value=row.category)
        ws.cell(row=r, column=13, value=row.kind)
        r += 1
    first = 2
    last = max(first, r - 1)  # με μηδέν γραμμές το range μένει ορθό, όχι G2:G1
    tr = last + 2
    ws.cell(row=tr, column=1, value="ΣΥΝΟΛΑ").font = bold
    for col in (7, 8, 9):
        L = get_column_letter(col)
        c = ws.cell(row=tr, column=col, value=f"=SUM({L}{first}:{L}{last})")
        c.number_format = "#,##0.00"; c.font = bold
    if name == "ΕΣΟΔΑ":
        c = ws.cell(row=tr, column=10, value=f"=SUM(J{first}:J{last})")
        c.number_format = "#,##0.00"; c.font = bold
        if subtotal_parties:
            for i, (afm, label) in enumerate(subtotal_parties, start=1):
                ws.cell(row=tr+i, column=1, value=f"ΣΥΝΟΛΑ {label} (ΑΦΜ {afm})").font = bold
                for col in (7, 8, 9, 10):
                    L = get_column_letter(col)
                    c = ws.cell(row=tr+i, column=col,
                                value=f'=SUMIF(F{first}:F{last},"{afm}",{L}{first}:{L}{last})')
                    c.number_format = "#,##0.00"; c.font = bold
            rest = tr + len(subtotal_parties) + 1
            ws.cell(row=rest, column=1, value="ΣΥΝΟΛΑ ΛΟΙΠΑ").font = bold
            for col in (7, 8, 9, 10):
                L = get_column_letter(col)
                minus = "".join(f"-{L}{tr+i}" for i in range(1, len(subtotal_parties) + 1))
                c = ws.cell(row=rest, column=col, value=f"={L}{tr}{minus}")
                c.number_format = "#,##0.00"; c.font = bold
    else:
        tr2 = tr + 1
        ws.cell(row=tr2, column=1, value="ΣΥΝΟΛΑ ΜΕ ΕΚΠΙΠΤΩΜΕΝΟ ΦΠΑ").font = bold
        c = ws.cell(row=tr2, column=7, value=f"=SUM(G{first}:G{last})")
        c.number_format = "#,##0.00"; c.font = bold
        c = ws.cell(row=tr2, column=8, value=f'=SUMIF(J{first}:J{last},"ΝΑΙ",H{first}:H{last})')
        c.number_format = "#,##0.00"; c.font = bold
    return len(rows)

def _normalize_nai_oxi(v):
    """Legacy αρχεία έχουν λατινικό NAI/OXI στη στήλη J — το SUMIF ταιριάζει μόνο ελληνικά."""
    if isinstance(v, str):
        s = v.strip().upper()
        if s in ("NAI", "ΝΑΙ"):
            return "ΝΑΙ"
        if s in ("OXI", "ΟΧΙ", "ΌΧΙ"):
            return "ΌΧΙ"
    return v

def read_sheet(path, name: str) -> list[SheetRow]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[name]
    out = []
    for row in ws.iter_rows(min_row=2):
        a, typ = row[0].value, row[1].value
        if isinstance(a, str) and a.strip().startswith("ΣΥΝΟΛΑ"):
            continue
        if typ is None or not str(typ).strip():
            continue
        d = row[4].value
        d = d.date() if isinstance(d, datetime) else d
        out.append(SheetRow(
            type=str(typ), series=str(row[2].value or ""), aa=str(row[3].value or ""),
            date=d, afm=str(row[5].value or ""),
            net=float(row[6].value or 0), vat=float(row[7].value or 0),
            total=float(row[8].value or 0), col_j=_normalize_nai_oxi(row[9].value),
            charact=str(row[10].value or ""), category=str(row[11].value or ""),
            kind=str(row[12].value or "")))
    return out
